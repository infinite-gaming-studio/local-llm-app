import subprocess
import tempfile
import os
import json


def extract_frames(video_path: str, max_frames: int = 5) -> list[str]:
    if not os.path.exists(video_path):
        return [f"Error: file not found: {video_path}"]

    frames = []
    try:
        result = subprocess.run(
            [
                "ffprobe", "-v", "error",
                "-show_entries", "format=duration",
                "-of", "json", video_path,
            ],
            capture_output=True, text=True, timeout=10,
        )
        info = json.loads(result.stdout)
        duration = float(info["format"]["duration"])
        interval = duration / (max_frames + 1)

        for i in range(max_frames):
            timestamp = interval * (i + 1)
            with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
                subprocess.run(
                    [
                        "ffmpeg", "-y",
                        "-ss", str(timestamp),
                        "-i", video_path,
                        "-vframes", "1",
                        "-q:v", "2",
                        tmp.name,
                    ],
                    capture_output=True, text=True, timeout=30,
                )
                with open(tmp.name, "rb") as f:
                    import base64
                    frames.append(base64.b64encode(f.read()).decode())
                os.unlink(tmp.name)

        return frames
    except Exception as e:
        return [f"Error extracting frames: {e}"]


def parse_document(path: str) -> str:
    if not os.path.exists(path):
        return f"Error: file not found: {path}"

    ext = os.path.splitext(path)[1].lower()

    if ext == ".txt":
        with open(path, "r", encoding="utf-8") as f:
            return f.read()

    ext_map = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".xlsx": "xlsx",
        ".csv": "csv",
    }
    format_type = ext_map.get(ext, ext)

    if format_type in ("pdf", "docx", "xlsx"):
        try:
            result = subprocess.run(
                ["python3", "-c", f"""
import sys
ext = '{format_type}'
path = '{path}'
if ext == 'pdf':
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        print('\\n'.join(p.page_text for p in pdf.pages[:20]))
elif ext == 'docx':
    from docx import Document
    doc = Document(path)
    print('\\n'.join(p.text for p in doc.paragraphs))
elif ext == 'xlsx':
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f'--- Sheet: {{sheet}} ---')
        for row in ws.iter_rows(values_only=True):
            print('\\t'.join(str(c) for c in row if c is not None))
"""],
                capture_output=True, text=True, timeout=30,
            )
            if result.returncode == 0:
                return result.stdout[:10000]
            return f"Error parsing document: {result.stderr}"
        except subprocess.TimeoutExpired:
            return "Error: document parsing timed out"

    if format_type == "csv":
        import csv
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = []
            for i, row in enumerate(reader):
                if i >= 50:
                    rows.append("... (truncated)")
                    break
                rows.append(", ".join(row))
            return "\n".join(rows)

    return f"Unsupported file type: {ext}"
